"""Render user separator drawings and extract supplier descriptions, without prices."""
import argparse
import hashlib
import json
import re
from pathlib import Path
import pypdfium2 as pdfium

PDFS=[('bdv60_5','сепаратор тип BDV60-5.pdf'),('fv8','Сепаратор тип FV8 16.09.2025-Разметка.pdf')]

def read_books(paths):
    """Only equipment descriptions, purposes and quantities enter the cache."""
    from openpyxl import load_workbook
    books=[]
    for path in paths:
        wb=load_workbook(path,read_only=True,data_only=True)
        ws=wb['S-4000'];rows=[]
        for cells in ws.iter_rows(min_row=1,min_col=9,max_col=11):
            row={cell.coordinate:cell.value for cell in cells if cell.value is not None}
            if row:rows.append(row)
        books.append(dict(file=path.name,sheet=ws.title,rows=rows));wb.close()
    return books

def inventory_sources(cache,directory,workbooks):
    from prepare_sources import SOURCES
    files=[directory/name for name in SOURCES.values()]+[directory/name for _,name in PDFS]+list(workbooks)
    records=[dict(name=p.name,path=str(p.resolve()),bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in files]
    path=cache/'source-inventory.json'
    if path.exists():
        old=json.loads(path.read_text(encoding='utf8'))
        assert {r['name']:r['sha256'] for r in old}=={r['name']:r['sha256'] for r in records},'Source files changed; use a new cache/version'
    else:path.write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf8')

def prepare(cache, directory, output,workbooks=None):
    cache.mkdir(parents=True,exist_ok=True)
    output.mkdir(parents=True, exist_ok=True)
    if workbooks:
        books=read_books(workbooks);inventory_sources(cache,directory,workbooks)
        (cache/'workbooks-equipment-only.json').write_text(json.dumps(books,ensure_ascii=False,indent=2),encoding='utf8')
    else:
        path=cache/'workbooks-equipment-only.json'
        if not path.exists():path=cache/'workbooks-raw.json'
        books=json.loads(path.read_text(encoding='utf8'))
    trims = {}
    for book, name in zip(books, ['standard', 'comfort', 'comfort_plus']):
        rows = []
        for source in book['rows']:
            for cell, value in source.items():
                if re.fullmatch(r'J\d+', cell) and isinstance(value, str):
                    n = int(cell[1:])
                    quantity = source.get('K' + str(n))
                    if isinstance(quantity, (int, float)) and n > 3:
                        rows.append(dict(row=n, description=value, quantity=quantity,
                                         purpose=source.get('I' + str(n), ''),
                                         cells=[cell, 'K' + str(n)]))
        trims[name] = dict(source=book['file'], sheet=book['sheet'], supplier='АДЛ', rows=rows)
    (output / 'trim-source-rows.json').write_text(json.dumps(trims, ensure_ascii=False, indent=2), encoding='utf8')
    if workbooks:
        proof=dict(status='DIRECT_WORKBOOK_EXTRACTION',workbooks=len(workbooks),sheet='S-4000',columns=['I','J','K'],
                   output_sha256=hashlib.sha256((output/'trim-source-rows.json').read_bytes()).hexdigest())
        (cache/'reference-checks.json').write_text(json.dumps(proof,indent=2),encoding='utf8')
    for stem, filename in PDFS:
        pdf = pdfium.PdfDocument(directory / filename)
        for n in range(len(pdf)):
            pdf[n].render(scale=2).to_pil().save(cache / (stem + '-page-' + str(n + 1) + '.png'))
        pdf.close()
        print('RENDERED', stem, flush=True)

if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('--cache', type=Path, required=True)
    p.add_argument('--directory', type=Path, required=True)
    p.add_argument('--output', type=Path, required=True)
    p.add_argument('--workbooks',type=Path,nargs=3,metavar=('STANDARD','COMFORT','COMFORT_PLUS'),help='Three source XLSX files in trim order')
    a = p.parse_args()
    prepare(a.cache, a.directory, a.output,a.workbooks)
