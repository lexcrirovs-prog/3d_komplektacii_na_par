"""Read equipment-only I/J/K cells from the three user-approved trim books."""
import hashlib
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

inventory, output = map(Path, sys.argv[1:])
sources = [s for s in json.loads(inventory.read_text(encoding='utf8')) if s['name'].endswith('.xlsx')]
assert len(sources) == 3
catalog = dict(date='2026-09-12', executor='Codex / GPT-6', trims={})
for trim, source in zip(['standard','comfort','comfort_plus'],sources):
    path = Path(source['path'])
    book = load_workbook(path,read_only=True,data_only=True)
    sheets = {}
    for sheet in book:
        if not sheet.title.strip().startswith('S-'): continue
        rows = []
        for cells in sheet.iter_rows(min_row=4,min_col=9,max_col=11):
            purpose, description, quantity = [c.value for c in cells]
            if isinstance(description,str) and isinstance(quantity,(int,float)):
                rows.append(dict(description=description, quantity=quantity,
                    purpose=purpose or '', row=cells[0].row))
        sheets[sheet.title.strip()] = rows
    catalog['trims'][trim] = dict(source=path.name,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),sheets=sheets)
    book.close()
output.parent.mkdir(parents=True,exist_ok=True)
output.write_text(json.dumps(catalog,ensure_ascii=False,indent=2),encoding='utf8')
print({k:{s:len(r) for s,r in v['sheets'].items()} for k,v in catalog['trims'].items()})
